# tests/test_specimen_and_metadata.py
import sys
import os
import math
import tempfile
import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from data_worker import DataWorker
from controllers import ExcelExportWorker
from execution.pdf_reporter import PDFReporter


def test_data_worker_sample_params_rectangular():
    worker = DataWorker(None)
    worker.update_sample_params({
        'sample_type': 'Rectangular',
        'width': 10.0,
        'thickness': 2.0,
        'gauge_length': 50.0
    })
    assert worker._sample_area == pytest.approx(20.0, rel=1e-5)
    assert worker._sample_gauge == 50.0


def test_data_worker_sample_params_circular_diameter():
    worker = DataWorker(None)
    worker.update_sample_params({
        'sample_type': 'Circular',
        'diameter': 10.0,
        'gauge_length': 50.0
    })
    expected_area = (math.pi / 4.0) * (10.0 ** 2)
    assert worker._sample_area == pytest.approx(expected_area, rel=1e-5)
    assert worker._sample_gauge == 50.0


def test_data_worker_sample_params_circular_radius_legacy():
    worker = DataWorker(None)
    worker.update_sample_params({
        'sample_type': 'Circular',
        'radius': 5.0,
        'gauge_length': 50.0
    })
    expected_area = math.pi * (5.0 ** 2)
    assert worker._sample_area == pytest.approx(expected_area, rel=1e-5)


def test_data_worker_sample_params_custom_area():
    worker = DataWorker(None)
    worker.update_sample_params({
        'sample_type': 'Custom / Direct Area',
        'custom_area': 42.75,
        'gauge_length': 75.0
    })
    assert worker._sample_area == pytest.approx(42.75, rel=1e-5)
    assert worker._sample_gauge == 75.0


def test_sample_panel_blank_defaults(qapp):
    from gui_panels import SamplePanel

    panel = SamplePanel()
    assert panel.width_inp.text() == ""
    assert panel.thick_inp.text() == ""
    assert panel.rect_gauge_inp.text() == ""
    assert panel.diam_inp.text() == ""
    assert panel.circ_gauge_inp.text() == ""
    assert panel.custom_area_inp.text() == ""
    assert panel.custom_gauge_inp.text() == ""
    assert "Calculated Area: —" in panel.lbl_calc_area.text()

    data = panel.get_sample_data()
    assert data["width"] is None
    assert data["thickness"] is None
    assert data["diameter"] is None
    assert data["radius"] is None
    assert data["custom_area"] is None
    assert data["gauge_length"] is None


def test_sample_panel_qt_gui(qapp):
    from gui_panels import SamplePanel

    panel = SamplePanel(
        init_sample_id='SPEC-01',
        init_sample_type='Circular',
        init_diam=8.0,
        init_custom_area=60.0,
        init_gauge=45.0,
        init_temp='700',
        init_stress='200',
        init_weight='30'
    )

    assert panel.sample_type_combo.currentText() == 'Circular'
    assert float(panel.diam_inp.text()) == 8.0
    assert 'Calculated Area: 50.27 mm' in panel.lbl_calc_area.text()

    data = panel.get_sample_data()
    assert data['sample_type'] == 'Circular'
    assert data['diameter'] == 8.0
    assert data['radius'] == 4.0
    assert data['temperature'] == '700'
    assert data['target_stress'] == '200'
    assert data['dead_weight'] == '30'

    panel.sample_type_combo.setCurrentText('Custom / Direct Area')
    assert panel.stack.currentIndex() == 2
    assert 'Calculated Area: 60.00 mm' in panel.lbl_calc_area.text()

    panel.custom_area_inp.setText('123.45')
    panel._update_calc_area()
    assert 'Calculated Area: 123.45 mm' in panel.lbl_calc_area.text()

    legacy_data = {
        'sample_id': 'LEGACY_SPEC',
        'sample_type': 'Circular',
        'radius': 6.0,
        'gauge_length': 60.0,
        'temperature': '650',
        'target_stress': '180',
        'dead_weight': '25'
    }
    panel.restore_sample_data(legacy_data)
    assert panel.sample_id_input.text() == 'LEGACY_SPEC'
    assert float(panel.diam_inp.text()) == 12.0
    assert panel.temp_inp.text() == '650'
    assert panel.target_stress_inp.text() == '180'
    assert panel.dead_weight_inp.text() == '25'


def test_excel_export_metadata_includes_new_fields():
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as tmpdir:
        meta = {
            'Sample ID': 'TEST-EXCEL',
            'Sample Type': 'Circular',
            'Gauge Length (mm)': 50.0,
            'Diameter (mm)': 10.0,
            'Cross-Sectional Area (mm2)': 78.54,
            'Test Temperature (C)': '650',
            'Target Nominal Stress (MPa)': '150',
            'Applied Dead Weight (kg)': '25.0'
        }
        worker = ExcelExportWorker(tmpdir, metadata=meta)
        worker.run()

        excel_path = os.path.join(tmpdir, 'test_results.xlsx')
        assert os.path.exists(excel_path)

        wb = load_workbook(excel_path)
        assert 'Sample Parameters' in wb.sheetnames
        ws = wb['Sample Parameters']
        rows = list(ws.iter_rows(values_only=True))
        param_dict = {r[0]: r[1] for r in rows if r[0] is not None}

        assert param_dict.get('Sample ID') == 'TEST-EXCEL'
        assert param_dict.get('Diameter (mm)') == '10.0'
        assert param_dict.get('Cross-Sectional Area (mm2)') == '78.54'
        assert param_dict.get('Test Temperature (C)') == '650'
        assert param_dict.get('Target Nominal Stress (MPa)') == '150'
        assert param_dict.get('Applied Dead Weight (kg)') == '25.0'


def test_pdf_reporter_with_new_metadata():
    from PIL import Image
    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path = os.path.join(tmpdir, 'test_report.pdf')
        dummy_graph = os.path.join(tmpdir, 'dummy_graph.png')
        img = Image.new('RGB', (100, 100), color=(73, 109, 137))
        img.save(dummy_graph)

        reporter = PDFReporter()
        meta = {
            'sample_id': 'PDF-SPEC-01',
            'sample_type': 'Circular',
            'diameter': 12.0,
            'area': 113.1,
            'gauge_length': 50.0,
            'temperature': '700',
            'target_stress': '250',
            'dead_weight': '35'
        }
        res = {
            'e_modulus': 200000,
            'yield_strength': 350,
            'uts': 500,
            'break_stress': 400
        }
        success = reporter.generate_report(pdf_path, meta, res, 'Specimen tested at 700C.', dummy_graph)
        assert success is True
        assert os.path.exists(pdf_path)
        assert os.path.getsize(pdf_path) > 1000


def test_dashboard_panel_graph_toggle(qapp):
    from gui_test_views import TestDashboardPanel

    panel = TestDashboardPanel()
    assert panel.btn_graph_ld.isChecked() is True
    assert panel.btn_graph_ss.isChecked() is False

    emitted = []
    panel.bl_graph_changed.connect(lambda mode: emitted.append(mode))

    # Switch to Stress vs Strain
    panel.btn_graph_ss.click()
    assert panel.btn_graph_ss.isChecked() is True
    assert panel.btn_graph_ld.isChecked() is False
    assert emitted[-1] == "Stress vs Strain"

    # Switch back to Load vs Displacement
    panel.btn_graph_ld.click()
    assert panel.btn_graph_ld.isChecked() is True
    assert panel.btn_graph_ss.isChecked() is False
    assert emitted[-1] == "Load vs Displacement"

    # Verify disable/enable
    panel.bl_graph_toggle.setEnabled(False)
    assert panel.bl_graph_toggle.isEnabled() is False
    assert panel.btn_graph_ld.isEnabled() is False
    assert panel.btn_graph_ss.isEnabled() is False

    panel.bl_graph_toggle.setEnabled(True)
    assert panel.bl_graph_toggle.isEnabled() is True

