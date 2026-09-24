# controllers.py
"""
Dedicated controllers for the Creep Testing Machine GUI.
"""
import os
import json
import time
import csv
import math

import pyqtgraph as pg

from PySide6.QtCore import QObject, QTimer, Signal, Slot, QThread
from PySide6.QtWidgets import QMessageBox, QProgressDialog

from typing import TYPE_CHECKING
if TYPE_CHECKING:
    from interfaces import SensorWorkerProtocol

from logging_config import get_logger, get_user_config_dir
from power_management import WindowsPowerManager
logger = get_logger(__name__)


# =========================================================
# ERROR MANAGER — Device error debounce
# =========================================================
class ErrorManager(QObject):
    """
    Collects device disconnection errors and shows a single consolidated
    popup after a 3-second debounce window.
    """

    def __init__(self, parent_widget):
        super().__init__(parent_widget)
        self._parent = parent_widget
        self._buffer = []
        self._last_flush_time = 0.0

        self._timer = QTimer(self)
        self._timer.setSingleShot(True)
        self._timer.setInterval(3000)
        self._timer.timeout.connect(self._flush)

    def collect(self, device_name, error_msg):
        """Buffer a device error. Deduplicates by device name."""
        if time.time() - self._last_flush_time < 5.0:
            logger.debug("%s error suppressed (recent flush): %s", device_name, error_msg)
            return
        existing = {e[0] for e in self._buffer}
        if device_name in existing:
            return
        self._buffer.append((device_name, error_msg))
        logger.error("%s error: %s", device_name, error_msg)
        if not self._timer.isActive():
            self._timer.start()

    def _flush(self):
        self._last_flush_time = time.time()
        errors = self._buffer.copy()
        self._buffer.clear()
        if not errors:
            return
        unique_names = list(dict.fromkeys(e[0] for e in errors))
        details = "\n\n".join(f"{n}:\n{m}" for n, m in errors)
        if len(unique_names) >= 2:
            mb = QMessageBox(self._parent)
            mb.setIcon(QMessageBox.Critical)
            mb.setWindowTitle("USB Hub Disconnected")
            mb.setText(f"Multiple devices lost connection.\nAffected: {', '.join(unique_names)}\n\nCheck the USB hub cable.")
            mb.setDetailedText(details)
            mb.setMinimumWidth(600)
            mb.exec()
        else:
            device_name, error_msg = errors[0]
            mb = QMessageBox(self._parent)
            mb.setIcon(QMessageBox.Critical)
            mb.setWindowTitle(f"{device_name} Disconnected")
            mb.setText(f"{device_name} lost connection.\n\nCheck the USB cable.")
            mb.setDetailedText(error_msg)
            mb.setMinimumWidth(500)
            mb.exec()


# =========================================================
# PRESET POSITION CONTROLLER — Multi-Preset Load / Unload
# =========================================================
class PresetPositionController(QObject):
    """Manages preset position (Multiple Named Load Presets / Unload) logic and config persistence."""

    motor_direction_request = Signal(int)    # 1 / -1 / 0
    presets_updated = Signal()

    POSITIONS_CONFIG_FILE = os.path.join(get_user_config_dir(), "preset_positions.json")
    MAX_LOAD_PRESETS = 10

    def __init__(self, serial_handler, uni_control, main_window):
        super().__init__(main_window)
        self._serial = serial_handler
        self._uni    = uni_control
        self._mw     = main_window
        if self._uni:
            self._uni.controller = self

        self._raw_encoder_pos      = 0

        self._load_presets         = []      # List of dicts: [{"id": str, "name": str, "pulses": int, "mm": float}]
        self._selected_load_id     = None
        self._unload_pos           = None
        
        self._travel_active        = False
        self._target_pos           = None
        self._target_name          = ""
        self._travel_rpm           = 0
        self._last_travel_cmd      = ""
        self._travel_initial_diff_sign = None

        self._timer = QTimer(self)
        self._timer.setInterval(20)
        self._timer.timeout.connect(self._travel_loop)

        self._load_config()
        self._sync_ui()

    def _load_config(self):
        try:
            if os.path.exists(self.POSITIONS_CONFIG_FILE):
                with open(self.POSITIONS_CONFIG_FILE, 'r') as f:
                    data = json.load(f)
                    
                    # 1. Load Presets (with legacy migration support)
                    raw_presets = data.get("load_positions")
                    if isinstance(raw_presets, list):
                        self._load_presets = raw_presets
                        for p in self._load_presets:
                            if "cm" not in p:
                                p["cm"] = round(self.pulses_to_cm(p.get("pulses", 0)), 4)
                            if "mm" not in p:
                                p["mm"] = round(self.pulses_to_mm(p.get("pulses", 0)), 4)
                    elif data.get("load_position") is not None:
                        # Migrate legacy single load position
                        legacy_p = int(data.get("load_position"))
                        self._load_presets = [{
                            "id": "load_default",
                            "name": "Default Load",
                            "pulses": legacy_p,
                            "cm": round(self.pulses_to_cm(legacy_p), 4),
                            "mm": round(self.pulses_to_mm(legacy_p), 4)
                        }]
                    else:
                        self._load_presets = []

                    # 2. Selected preset ID
                    self._selected_load_id = data.get("selected_load_id")
                    if self._load_presets:
                        valid_ids = [p["id"] for p in self._load_presets]
                        if self._selected_load_id not in valid_ids:
                            self._selected_load_id = valid_ids[0]
                    else:
                        self._selected_load_id = None

                    # 3. Unload position
                    self._unload_pos = data.get("unload_position")
                    logger.info("Loaded presets: %d load presets, Unload=%s", len(self._load_presets), self._unload_pos)
        except (json.JSONDecodeError, IOError) as e:
            logger.warning("Could not load preset config: %s", e)
            self._load_presets = []
            self._selected_load_id = None
            self._unload_pos = None

    def _save_config(self):
        try:
            active_p = self.get_selected_load_preset()
            legacy_load_pos = active_p["pulses"] if active_p else None
            with open(self.POSITIONS_CONFIG_FILE, 'w') as f:
                json.dump({
                    "load_positions": self._load_presets,
                    "selected_load_id": self._selected_load_id,
                    "unload_position": self._unload_pos,
                    "load_position": legacy_load_pos
                }, f, indent=2)
            logger.info("Saved preset positions (%d load presets).", len(self._load_presets))
        except IOError as e:
            logger.error("Could not save preset config: %s", e)

    def _sync_ui(self):
        if hasattr(self._uni, "update_presets_ui"):
            self._uni.update_presets_ui(
                load_presets=self._load_presets,
                selected_id=self._selected_load_id,
                unload_pos=self._unload_pos
            )

    @property
    def current_position(self):
        return self._raw_encoder_pos

    @staticmethod
    def pulses_to_cm(pulses):
        """Convert encoder pulses to cm: (pulses/10000) * 0.2."""
        return (pulses / 10000.0) * 0.2

    @staticmethod
    def cm_to_pulses(cm):
        """Convert cm to encoder pulses: (cm / 0.2) * 10000."""
        return int(round((cm / 0.2) * 10000.0))

    @staticmethod
    def pulses_to_mm(pulses):
        """Convert encoder pulses to mm: (pulses/10000) * 2.0."""
        return (pulses / 10000.0) * 2.0

    @staticmethod
    def mm_to_pulses(mm):
        """Convert mm to encoder pulses: (mm / 2.0) * 10000."""
        return int(round((mm / 2.0) * 10000.0))

    @Slot(int)
    def on_encoder_updated(self, raw_pos):
        self._raw_encoder_pos = raw_pos

    # --- Multiple Load Presets API ---
    def get_load_presets(self):
        return [dict(p) for p in self._load_presets]

    def get_selected_load_id(self):
        return self._selected_load_id

    def get_selected_load_preset(self):
        for p in self._load_presets:
            if p["id"] == self._selected_load_id:
                return p
        if self._load_presets:
            return self._load_presets[0]
        return None

    def set_selected_load_id(self, preset_id):
        if self._selected_load_id != preset_id:
            self._selected_load_id = preset_id
            self._save_config()
            self._sync_ui()

    def add_load_preset(self, name: str, pulses: int) -> tuple[bool, str]:
        if len(self._load_presets) >= self.MAX_LOAD_PRESETS:
            return False, f"Maximum limit of {self.MAX_LOAD_PRESETS} presets reached. Please delete an unused preset first."
        
        clean_name = name.strip() if name else ""
        if not clean_name:
            clean_name = f"Load Pos {len(self._load_presets) + 1}"
            
        preset_id = f"pos_{int(time.time() * 1000)}"
        cm_val = round(self.pulses_to_cm(pulses), 4)
        mm_val = round(self.pulses_to_mm(pulses), 4)
        new_preset = {
            "id": preset_id,
            "name": clean_name,
            "pulses": int(pulses),
            "cm": cm_val,
            "mm": mm_val
        }
        self._load_presets.append(new_preset)
        self._selected_load_id = preset_id
        self._save_config()
        self._sync_ui()
        self.presets_updated.emit()
        return True, "Preset added successfully."

    def delete_load_preset(self, preset_id: str) -> bool:
        initial_len = len(self._load_presets)
        self._load_presets = [p for p in self._load_presets if p["id"] != preset_id]
        if len(self._load_presets) < initial_len:
            if self._selected_load_id == preset_id:
                self._selected_load_id = self._load_presets[0]["id"] if self._load_presets else None
            self._save_config()
            self._sync_ui()
            self.presets_updated.emit()
            return True
        return False

    def rename_load_preset(self, preset_id: str, new_name: str) -> bool:
        clean_name = new_name.strip() if new_name else ""
        if not clean_name:
            return False
        for p in self._load_presets:
            if p["id"] == preset_id:
                p["name"] = clean_name
                self._save_config()
                self._sync_ui()
                self.presets_updated.emit()
                return True
        return False

    # --- Unload Position API ---
    def set_unload_pos(self):
        self._unload_pos = self.current_position
        self._save_config()
        self._sync_ui()
        cm = self.pulses_to_cm(self._unload_pos)
        QMessageBox.information(self._mw, "Unload Position Set", f"Unload position saved!\n{cm:.2f} cm ({self._unload_pos:,} pulses)")

    def clear_unload_pos(self):
        if self._unload_pos is None:
            QMessageBox.information(self._mw, "No Unload Position", "No unload position is currently set.")
            return
        reply = QMessageBox.question(
            self._mw, "Clear Unload Position", "Clear the saved unload position?",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        if reply == QMessageBox.Yes:
            self._unload_pos = None
            self._save_config()
            self._sync_ui()
            logger.info("Unload position CLEARED")

    # --- Travel Logic ---
    def start_travel_to(self, target_type, rpm):
        if target_type == "load":
            preset = self.get_selected_load_preset()
            if not preset:
                QMessageBox.warning(self._mw, "No Load Position", "No load position is selected. Use 'Manage Load Positions' to configure positions.")
                return
            target = preset["pulses"]
            name = preset["name"]
        else:
            target = self._unload_pos
            name = "Unload"

        if target is None:
            QMessageBox.warning(self._mw, f"No {name} Position", f"Set an {name} position first.")
            return

        target_cm = self.pulses_to_cm(target)
        reply = QMessageBox.question(
            self._mw, f"Confirm Travel to {name}",
            f"Are you sure you want to travel to '{name}'?\n\nTarget: {target_cm:.2f} cm ({target:,} pulses)\n\nPlease ensure the crosshead path is clear.",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        if reply != QMessageBox.Yes:
            return

        if not self._serial.is_connected():
            QMessageBox.critical(self._mw, "Connection Error", "USB is disconnected!")
            return

        self._travel_active = True
        self._target_pos = target
        self._target_name = name
        self._travel_rpm = abs(rpm)
        self._last_travel_cmd = ""
        self._travel_initial_diff_sign = None

        if hasattr(self._uni, "set_travel_active_state"):
            self._uni.set_travel_active_state(True)

        self._timer.start()
        logger.info("Travel STARTED to %s — target: %d, speed: %d RPM", name, target, self._travel_rpm)

    def stop_travel(self):
        self._travel_active = False
        self._timer.stop()
        self._serial.send_cmd("STOP")
        self.motor_direction_request.emit(0)
        if hasattr(self._uni, "set_travel_active_state"):
            self._uni.set_travel_active_state(False)
        logger.info("Travel STOPPED")

    def _travel_loop(self):
        if not self._travel_active:
            return
        curr   = self.current_position
        target = self._target_pos
        diff   = curr - target
        curr_cm = self.pulses_to_cm(curr)
        if hasattr(self._uni, "lbl_live_encoder"):
            self._uni.lbl_live_encoder.setText(f"Live: {curr_cm:.2f} cm ({curr:,} pulses)")

        # Deadband: 5000 pulses
        if abs(diff) <= 5000:
            self.stop_travel()
            QMessageBox.information(self._mw, "Travel Complete", f"Reached {self._target_name} position!\n{curr_cm:.2f} cm ({curr:,} pulses)")
            return

        if curr < target:
            cmd, dir_val = f"VEL:{self._travel_rpm}", 1
        else:
            cmd, dir_val = f"VEL:-{self._travel_rpm}", -1

        if cmd != self._last_travel_cmd:
            self._serial.send_cmd(cmd)
            self._serial.send_cmd("START")
            self.motor_direction_request.emit(dir_val)
            self._last_travel_cmd = cmd
            self._travel_initial_diff_sign = (diff > 0)

        if self._travel_initial_diff_sign is not None:
            if (diff > 0) != self._travel_initial_diff_sign:
                self.stop_travel()
                if abs(diff) <= 5000:
                    QMessageBox.information(self._mw, "Travel Complete", f"Reached {self._target_name} position!\n{curr_cm:.2f} cm ({curr:,} pulses)")
                else:
                    overshoot_cm = self.pulses_to_cm(abs(diff))
                    QMessageBox.warning(
                        self._mw, "Travel Overshoot",
                        f"Overshot {self._target_name} position by {overshoot_cm:.2f} cm ({abs(diff):,} pulses)!\nThis exceeds the 5000-pulse deadband.\nTry a lower speed."
                    )



# =========================================================
# EXCEL EXPORT WORKER (Background Thread)
# =========================================================
class ExcelExportWorker(QObject):
    """Runs CSV-to-Excel conversion on a background thread."""
    finished = Signal()
    error    = Signal(str)

    def __init__(self, save_dir, metadata=None, sample_id=None):
        super().__init__()
        self._save_dir = save_dir
        if isinstance(metadata, dict):
            self._metadata = dict(metadata)
            if sample_id and "Sample ID" not in self._metadata:
                self._metadata["Sample ID"] = sample_id
        elif isinstance(metadata, str):
            self._metadata = {"Sample ID": metadata}
        elif sample_id:
            self._metadata = {"Sample ID": sample_id}
        else:
            self._metadata = {}

    @Slot()
    def run(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            excel_path = os.path.join(self._save_dir, "test_results.xlsx")
            csv_path   = os.path.join(self._save_dir, "test_result(csv format).csv")
            wb = Workbook()
            ws = wb.active
            ws.title = "Sensor Data"
            row_count = 0
            if os.path.exists(csv_path):
                with open(csv_path, "r", encoding="utf-8") as f:
                    for row in csv.reader(f):
                        ws.append(row)
                        row_count += 1
            if row_count > 0:
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                ws.column_dimensions['A'].width = 16
                ws.column_dimensions['B'].width = 16
                ws.column_dimensions['C'].width = 16

            ws_p = wb.create_sheet("Sample Parameters")
            ws_p.append(["Parameter", "Value"])
            for k, v in self._metadata.items():
                ws_p.append([k, str(v)])
            for cell in ws_p[1]:
                cell.font = Font(bold=True)
            ws_p.column_dimensions['A'].width = 32
            ws_p.column_dimensions['B'].width = 24

            wb.save(excel_path)
            self.finished.emit()
        except Exception as e:
            self.error.emit(str(e))


# =========================================================
# RECORDING CONTROLLER
# =========================================================
class RecordingController(QObject):
    recording_started = Signal()
    recording_stopped = Signal()

    def __init__(self, main_window, data_worker, camera_panel):
        super().__init__(main_window)
        self._mw          = main_window
        self._data_worker = data_worker
        self._camera      = camera_panel
        self.current_save_dir   = ""
        self.manual_captures_dir = None
        self._is_recording = False
        self._export_thread = None
        self._export_worker = None

    def start_record(self):
        mw = self._mw
        if not mw.serial_handler.is_connected():
            QMessageBox.critical(mw, "Connection Error", "USB is disconnected! Please connect to the device.")
            return
        if mw._save_images and not self._camera.is_live():
            QMessageBox.critical(mw, "Camera Error", "Camera is not live. Start the stream in Image Setup first.")
            return
        if not getattr(mw, "_output_folder", None):
            QMessageBox.critical(mw, "No Output Folder", "Please select an output folder in Test Configuration.")
            return
        if not getattr(mw, "_save_sensor_data", False) and not getattr(mw, "_save_images", False):
            QMessageBox.critical(mw, "No Save Options", "Please select at least one save option.")
            return

        timestamp = time.strftime("%Y%m%d_%H%M%S")
        test_folder = f"{getattr(mw, '_sample_id', 'Test')}_{timestamp}"
        self.current_save_dir = os.path.join(mw._output_folder, test_folder)
        try:
            os.makedirs(self.current_save_dir, exist_ok=True)
            self.manual_captures_dir = os.path.join(self.current_save_dir, "Manual_Image_Captures")
            os.makedirs(self.manual_captures_dir, exist_ok=True)
            if mw._save_images:
                self.images_dir = os.path.join(self.current_save_dir, "Auto_Record_Image_Captures")
                os.makedirs(self.images_dir, exist_ok=True)
            else:
                self.images_dir = None
        except Exception as e:
            QMessageBox.critical(mw, "Folder Error", f"Failed to create test folder:\n{e}")
            return

        if mw._save_images and self.images_dir:
            self._camera.worker.start_recording(self.images_dir)

        csv_path = ""
        if getattr(mw, "_save_sensor_data", False):
            csv_path = os.path.join(self.current_save_dir, "test_result(csv format).csv")
            img_dir  = self.images_dir if self.images_dir and mw._save_images else ""
            rate = 20
            if hasattr(mw, "sensor_panel") and hasattr(mw.sensor_panel, "get_sampling_rate"):
                rate = mw.sensor_panel.get_sampling_rate()
            sample_data = mw.sample_panel.get_sample_data() if hasattr(mw, "sample_panel") else {}
            output_folder = getattr(mw, "_output_folder", "")
            self._data_worker.request_start_recording(csv_path, img_dir, rate, output_folder, sample_data)

        if hasattr(mw, "sensor_panel") and hasattr(mw.sensor_panel, "set_recording_state"):
            mw.sensor_panel.set_recording_state(True)
        if hasattr(mw, "dashboard_panel"):
            if hasattr(mw.dashboard_panel, "bl_graph_toggle"):
                mw.dashboard_panel.bl_graph_toggle.setEnabled(False)
            elif hasattr(mw.dashboard_panel, "bl_graph_combo"):
                mw.dashboard_panel.bl_graph_combo.setEnabled(False)
            mw.dashboard_panel.start_rec_timer()

        self._is_recording = True
        WindowsPowerManager.prevent_sleep()
        self.recording_started.emit()

    def stop_record(self):
        mw = self._mw
        if getattr(mw, "_save_images", False):
            self._camera.worker.stop_recording()
        self._data_worker.stop_recording()
        if hasattr(mw, "sensor_panel") and hasattr(mw.sensor_panel, "set_recording_state"):
            mw.sensor_panel.set_recording_state(False)
        if hasattr(mw, "dashboard_panel"):
            if hasattr(mw.dashboard_panel, "bl_graph_toggle"):
                mw.dashboard_panel.bl_graph_toggle.setEnabled(True)
            elif hasattr(mw.dashboard_panel, "bl_graph_combo"):
                mw.dashboard_panel.bl_graph_combo.setEnabled(True)
            mw.dashboard_panel.stop_rec_timer()

        if getattr(mw, "_save_sensor_data", False):
            self._run_excel_export()
        else:
            self._show_save_complete()

        self._is_recording = False
        self.manual_captures_dir = None
        WindowsPowerManager.allow_sleep()
        self.recording_stopped.emit()

    def _run_excel_export(self):
        mw = self._mw
        self._export_progress = QProgressDialog("Exporting data to Excel...", None, 0, 0, mw)
        self._export_progress.setWindowTitle("Saving")
        self._export_progress.setMinimumWidth(320)
        self._export_progress.setMinimumDuration(0)
        self._export_progress.setCancelButton(None)
        self._export_progress.setModal(True)
        self._export_progress.show()

        gl = getattr(mw, "_sample_gauge_length", None)
        metadata = {
            "Sample ID": getattr(mw, "_sample_id", "Test"),
            "Sample Type": getattr(mw, "_sample_type", "Rectangular"),
            "Gauge Length (mm)": gl if gl is not None else "--",
            "Sampling Rate (samples/hour)": getattr(self._data_worker, "_sampling_rate", 20),
            "Tare Load Offset (N)": round(getattr(self._data_worker, "_load_offset", 0.0), 3),
            "Tare Disp Offset (mm)": round(getattr(self._data_worker, "_disp_offset", 0.0), 4),
            "Max Displacement (mm)": round(getattr(self._data_worker, "_max_disp", 0.0), 4),
            "Export Timestamp": time.strftime("%Y-%m-%d %H:%M:%S")
        }
        temp = getattr(mw, "_sample_temperature", "")
        if temp:
            metadata["Test Temperature (°C)"] = temp
        stress = getattr(mw, "_sample_target_stress", "")
        if stress:
            metadata["Target Nominal Stress (MPa)"] = stress
        weight = getattr(mw, "_sample_dead_weight", "")
        if weight:
            metadata["Applied Dead Weight (kg)"] = weight

        stype = getattr(mw, "_sample_type", "Rectangular")
        if stype == "Circular":
            diam = getattr(mw, "_sample_diameter", None)
            if diam is None and getattr(mw, "_sample_radius", None) is not None:
                diam = getattr(mw, "_sample_radius") * 2.0
            if diam is not None:
                area = (math.pi / 4.0) * (diam ** 2)
                metadata["Diameter (mm)"] = diam
                metadata["Cross-Sectional Area (mm²)"] = round(area, 3)
            else:
                metadata["Diameter (mm)"] = "--"
                metadata["Cross-Sectional Area (mm²)"] = "--"
        elif stype == "Custom / Direct Area":
            area = getattr(mw, "_sample_custom_area", None)
            if area is not None:
                metadata["Cross-Sectional Area (mm²)"] = round(area, 3)
            else:
                metadata["Cross-Sectional Area (mm²)"] = "--"
        else:
            w = getattr(mw, "_sample_width", None)
            th = getattr(mw, "_sample_thickness", None)
            if w is not None and th is not None:
                area = w * th
                metadata["Width (mm)"] = w
                metadata["Thickness (mm)"] = th
                metadata["Cross-Sectional Area (mm²)"] = round(area, 3)
            else:
                metadata["Width (mm)"] = w if w is not None else "--"
                metadata["Thickness (mm)"] = th if th is not None else "--"
                metadata["Cross-Sectional Area (mm²)"] = "--"

        self._export_thread = QThread()
        self._export_worker = ExcelExportWorker(
            self.current_save_dir,
            metadata
        )
        self._export_worker.moveToThread(self._export_thread)
        self._export_thread.started.connect(self._export_worker.run)
        self._export_worker.finished.connect(self._on_export_finished)
        self._export_worker.error.connect(self._on_export_error)
        self._export_worker.finished.connect(self._export_thread.quit)
        self._export_worker.error.connect(self._export_thread.quit)
        self._export_thread.start()

    def _on_export_finished(self):
        self._export_progress.close()
        self._show_save_complete(excel_success=True)
        self._export_thread.wait()

    def _on_export_error(self, error_msg):
        self._export_progress.close()
        logger.error("Excel export error: %s", error_msg)
        if "openpyxl" in error_msg.lower():
            mw = self._mw
            saved = []
            if getattr(mw, "_save_sensor_data", False): saved.append("Sensor data (CSV format)")
            if getattr(mw, "_save_images",      False): saved.append("Camera images")
            QMessageBox.warning(mw, "Excel Export Skipped",
                f"{' and '.join(saved)} saved to:\n\n{self.current_save_dir}\n\n"
                "Note: Excel (.xlsx) file was not generated because 'openpyxl' is not installed.\n"
                "To enable automatic .xlsx generation, run in terminal:\n\n"
                "pip install openpyxl")
        else:
            self._show_save_complete(excel_success=False)
        self._export_thread.wait()

    def _show_save_complete(self, excel_success=True):
        mw = self._mw
        saved = []
        if getattr(mw, "_save_sensor_data", False):
            saved.append("Sensor data (Excel & CSV)" if excel_success else "Sensor data (CSV)")
        if getattr(mw, "_save_images",      False):
            saved.append("Camera images")
        QMessageBox.information(mw, "Recording Saved",
            f"{' and '.join(saved)} saved to:\n\n{self.current_save_dir}")


# =========================================================
# CAPTURE CONTROLLER — manual single-image capture
# =========================================================
class CaptureController(QObject):
    """Handles manual image capture."""
    capture_success = Signal(str)
    capture_error   = Signal(str)

    def __init__(self, main_window, camera_panel, dashboard_panel=None):
        super().__init__(main_window)
        self._mw        = main_window
        self._camera    = camera_panel
        self._dashboard = dashboard_panel

    def capture_image(self):
        mw = self._mw
        if not self._camera.is_live():
            QMessageBox.warning(mw, "Camera Not Active",
                "Camera stream is not active.\nStart the stream in 'Image Setup'.")
            self.capture_error.emit("Stream not active")
            return
        if not mw._output_folder:
            QMessageBox.warning(mw, "No Save Location",
                "No output folder selected.\nSet one in 'Test Configuration'.")
            self.capture_error.emit("No save location")
            return
        if not os.path.isdir(mw._output_folder):
            QMessageBox.warning(mw, "Save Location Invalid",
                "The output folder no longer exists.\nSelect a new one.")
            self.capture_error.emit("Folder missing")
            return

        if hasattr(mw, 'rec_ctrl') and mw.rec_ctrl._is_recording and mw.rec_ctrl.manual_captures_dir:
            captures_dir = mw.rec_ctrl.manual_captures_dir
        else:
            captures_dir = os.path.join(mw._output_folder, "Manual_Image_Captures")
            try:
                os.makedirs(captures_dir, exist_ok=True)
            except Exception as e:
                QMessageBox.critical(mw, "Folder Error", f"Failed to create captures folder:\n{e}")
                self.capture_error.emit("Folder creation failed")
                return

        timestamp  = time.strftime("%Y%m%d_%H%M%S")
        micro      = f"{time.time() % 1:.6f}"[2:]
        save_path  = os.path.join(captures_dir, f"capture_{timestamp}_{micro}.tiff")

        if hasattr(self._camera, 'btn_capture'):
            self._camera.btn_capture.setEnabled(False)
            from PySide6.QtCore import QTimer
            QTimer.singleShot(200, lambda: self._camera.btn_capture.setEnabled(True))
        if self._dashboard and hasattr(self._dashboard, 'btn_capture_dash'):
            self._dashboard.btn_capture_dash.setEnabled(False)
            from PySide6.QtCore import QTimer
            QTimer.singleShot(200, lambda: self._dashboard.btn_capture_dash.setEnabled(True))

        self._camera.worker.request_capture(save_path)
        logger.info("Manual capture requested: %s", save_path)


# =========================================================
# GRAPH CONTROLLER — dashboard curves & graph type switching
# =========================================================
class GraphController(QObject):
    """Manages dashboard graph curves for creep: Load(N) vs Time, Disp(mm) vs Time."""

    def __init__(self, main_window, dashboard_panel):
        super().__init__(main_window)
        self._mw   = main_window
        self._dash = dashboard_panel

        self._bl_graph_mode = "Load vs Displacement"

        self.load_t_curve  = None   # Load vs Time
        self.disp_t_curve  = None   # Disp vs Time
        self.load_d_curve  = None   # Load vs Disp

    def setup_curves(self, mode="Creep"):
        """Create dashboard plot curves."""
        self._dash.plot_load_time.clear()
        self._dash.plot_load_disp.clear()
        self._dash.plot_disp_time.clear()

        self.load_t_curve = self._dash.plot_load_time.plot(
            pen=pg.mkPen(color='#2980b9', width=3),
            autoDownsample=True, clipToView=True)
        self.load_d_curve = self._dash.plot_load_disp.plot(
            pen=pg.mkPen(color='#27ae60', width=3),
            autoDownsample=True, clipToView=True)
        self.disp_t_curve = self._dash.plot_disp_time.plot(
            pen=pg.mkPen(color='#e67e22', width=3),
            autoDownsample=True, clipToView=True)

    def on_bl_graph_changed(self, mode):
        self._bl_graph_mode = mode

    @Slot(dict)
    def update_graphs(self, data):
        """Route creep plot data to dashboard curves."""
        t     = data.get("time", [])
        load  = data.get("load", [])
        disp  = data.get("disp", [])

        if self.load_t_curve is not None and t:
            self.load_t_curve.setData(t, load)

        if self.disp_t_curve is not None and t:
            self.disp_t_curve.setData(t, disp)

        if self.load_d_curve is not None and disp:
            if self._bl_graph_mode == "Stress vs Strain":
                stress = data.get("stress", [])
                strain = data.get("strain", [])
                if stress and strain:
                    self.load_d_curve.setData(strain, stress)
                else:
                    self.load_d_curve.setData(disp, load)
            else:
                self.load_d_curve.setData(disp, load)
